"""Lecture conservatrice de CSV/XLSX d'export ORLI, sans connexion au SI.

Le schéma réel d'ORLI reste à faire confirmer : un mapping explicite est nécessaire.
Les cellules/formules sont traitées comme des données, jamais exécutées.
"""
from __future__ import annotations
import csv
import io
import json
import zipfile
from collections import Counter
from hashlib import sha256
from openpyxl import load_workbook
from fastapi import HTTPException

MAX_BYTES = 3 * 1024 * 1024
MAX_ROWS = 1000
FIELDS = ('orli_ref', 'order_ref', 'client_name', 'designation', 'quantity', 'status')
LIMITS = {'orli_ref':100, 'order_ref':100, 'client_name':150, 'designation':180, 'quantity':60, 'status':80}


def trim(value):
    if value is None: return ''
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def normalize_reference(value):
    # Aucune normalisation destructrice (pas d'effacement de zéros, tirets, etc.).
    return ' '.join(trim(value).split()).casefold()


def parse_mapping(text: str) -> dict[str, str]:
    try: obj = json.loads(text)
    except (ValueError, TypeError): raise HTTPException(422, 'Mapping JSON invalide')
    if not isinstance(obj, dict) or set(obj) - set(FIELDS) or 'orli_ref' not in obj:
        raise HTTPException(422, 'Mapping attendu : objet avec au moins orli_ref et des champs reconnus')
    if any(not isinstance(v, str) or not v.strip() or len(v) > 120 for v in obj.values()):
        raise HTTPException(422, 'Les noms de colonnes du mapping doivent être renseignés')
    if len(set(obj.values())) != len(obj):
        raise HTTPException(422, 'Une colonne ne peut alimenter deux champs métier')
    return {k: v.strip() for k, v in obj.items()}


def read_export(data: bytes, filename: str, mapping: dict[str, str]):
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext not in {'csv','xlsx'}: raise HTTPException(422, 'Seuls les exports CSV ou XLSX sont acceptés')
    if not data or len(data) > MAX_BYTES: raise HTTPException(413, 'Export vide ou supérieur à 3 Mo')
    if ext == 'csv':
        try: content = data.decode('utf-8-sig')
        except UnicodeDecodeError: raise HTTPException(422, 'CSV attendu en UTF-8 (avec ou sans BOM)')
        try: dialect=csv.Sniffer().sniff(content[:4096], delimiters=',;\t|')
        except csv.Error: dialect=csv.excel
        iterator = csv.reader(io.StringIO(content), dialect=dialect)
        try: headers = [trim(x) for x in next(iterator)]
        except StopIteration: raise HTTPException(422, 'Fichier sans en-tête')
        rows = list(iterator)
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos=archive.infolist()
                if len(infos)>3000 or sum(x.file_size for x in infos)>30*1024*1024 or any(x.file_size>20*1024*1024 for x in infos):
                    raise HTTPException(413, 'Classeur trop volumineux une fois décompressé')
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                ws=workbook.worksheets[0]
                iterator=ws.iter_rows(values_only=True)
                headers=[trim(x) for x in next(iterator)]
                rows=[]
                for i, row in enumerate(iterator):
                    if i > MAX_ROWS: raise HTTPException(413, 'Export limité à 1000 lignes')
                    rows.append(row)
            finally: workbook.close()
        except HTTPException: raise
        except Exception: raise HTTPException(422, 'Classeur XLSX illisible ou invalide')
    if not headers or len(headers)>80 or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise HTTPException(422, 'En-têtes vides, dupliqués ou trop nombreux')
    if any(col not in headers for col in mapping.values()):
        raise HTTPException(422, 'Colonnes du mapping absentes : '+', '.join(col for col in mapping.values() if col not in headers))
    if len(rows)>MAX_ROWS: raise HTTPException(413, 'Export limité à 1000 lignes')
    positions={k:headers.index(col) for k,col in mapping.items()}
    prepared=[]
    for line, row in enumerate(rows, start=2):
        if not any(trim(x) for x in row): continue
        values={k:trim(row[pos]) if pos < len(row) else '' for k,pos in positions.items()}
        errors=[]
        if len(row)>len(headers) and any(trim(x) for x in row[len(headers):]): errors.append('Colonnes supplémentaires non déclarées')
        if not values['orli_ref']: errors.append('Référence de lancement ORLI manquante')
        for k, v in values.items():
            if len(v)>LIMITS[k]: errors.append(f'{k} dépasse {LIMITS[k]} caractères')
            if ext=='xlsx' and v.startswith('='): errors.append(f'{k} : formule Excel non autorisée')
        # Ne conserver aucune valeur trop longue en base.
        normalized={k:(v[:LIMITS[k]] or None) for k,v in values.items()}
        prepared.append({'row_number':line,'fields':normalized,'errors':errors})
    counts=Counter(normalize_reference(x['fields']['orli_ref']) for x in prepared if x['fields']['orli_ref'])
    for row in prepared:
        key=normalize_reference(row['fields']['orli_ref'])
        if key and counts[key]>1:
            row['errors'].append('Référence ORLI répétée dans cet export')
    return {'format':ext, 'sha256':sha256(data).hexdigest(), 'headers':headers,
            'rows':prepared,'total_rows':len(prepared),'valid_rows':sum(not x['errors'] for x in prepared)}
